from pptx import Presentation

P = r"c:\Users\yuanx\CodeBuddy\20260922075003\在线购物系统_阶段汇报_五人分工.pptx"
EMU = 914400
W, H = 10.0, 5.625

prs = Presentation(P)
issues = []
for i, slide in enumerate(prs.slides, 1):
    for sh in slide.shapes:
        if sh.left is None or sh.top is None:
            continue
        l = sh.left / EMU
        t = sh.top / EMU
        w = (sh.width or 0) / EMU
        h = (sh.height or 0) / EMU
        r, b = l + w, t + h
        over = []
        if l < -0.02:
            over.append(f"left={l:.2f}")
        if t < -0.02:
            over.append(f"top={t:.2f}")
        if r > W + 0.02:
            over.append(f"right={r:.2f}")
        if b > H + 0.02:
            over.append(f"bottom={b:.2f}")
        if over:
            txt = ""
            if sh.has_text_frame:
                txt = sh.text_frame.text.replace("\n", " ")[:38]
            issues.append(f"P{i:02d} {sh.shape_type} '{sh.name}' [{', '.join(over)}] :: {txt}")

print(f"total slides: {len(prs.slides.__iter__.__self__._sldIdLst)}")
if issues:
    print("OUT-OF-BOUNDS:")
    for x in issues:
        print(" ", x)
else:
    print("no out-of-bounds shapes")
